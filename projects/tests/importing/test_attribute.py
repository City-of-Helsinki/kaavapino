import pytest
from openpyxl import Workbook

from projects.models import (
    Attribute,
    AttributeAutoValue,
    AttributeAutoValueMapping,
    AttributeValueChoice,
    CommonProjectPhase,
    Deadline,
    DocumentLinkFieldSet,
    DocumentLinkSection,
    FieldSetAttribute,
    ProjectPhaseDeadlineSection,
    ProjectPhaseDeadlineSectionAttribute,
    ProjectPhase,
    ProjectPhaseFieldSetAttributeIndex,
    ProjectPhaseSection,
    ProjectPhaseSectionAttribute,
    ProjectCardSection,
    ProjectCardSectionAttribute,
    ProjectFloorAreaSection,
    ProjectFloorAreaSectionAttribute,
    ProjectFloorAreaSectionAttributeMatrixCell,
    ProjectFloorAreaSectionAttributeMatrixStructure,
    ProjectSubtype,
)
from projects.models.attribute import AttributeCategorization
from projects.importing import AttributeImporter
from projects.importing.attribute import (
    ATTRIBUTE_API_VISIBILITY,
    ATTRIBUTE_ASSISTIVE_TEXT,
    ATTRIBUTE_CATEGORIZATION_COLUMNS,
    ATTRIBUTE_AUTO_VALUE_KEY_FIELD,
    ATTRIBUTE_AUTO_VALUE_MAPPING,
    ATTRIBUTE_BROADCAST_CHANGES,
    ATTRIBUTE_CHARACTER_LIMIT,
    ATTRIBUTE_CHOICES_REF,
    ATTRIBUTE_DATA_RETENTION,
    ATTRIBUTE_DEADLINE_SECTION_COLUMNS,
    ATTRIBUTE_FIELDSET,
    ATTRIBUTE_FLOOR_AREA_SECTION,
    ATTRIBUTE_FLOOR_AREA_SECTION_MATRIX_ROW,
    ATTRIBUTE_FLOOR_AREA_SECTION_MATRIX_CELL,
    ATTRIBUTE_EDIT_PRIVILEGE,
    ATTRIBUTE_ERROR,
    ATTRIBUTE_ERROR_TEXT,
    ATTRIBUTE_FIELD_ROLE,
    ATTRIBUTE_FIELD_SUBROLE,
    ATTRIBUTE_FIELDSET_TOTAL,
    ATTRIBUTE_GROUP,
    ATTRIBUTE_HIGHLIGHT_GROUP,
    ATTRIBUTE_IDENTIFIER,
    ATTRIBUTE_LINKED_FIELDS,
    ATTRIBUTE_MULTIPLE_CHOICE,
    ATTRIBUTE_NAME,
    ATTRIBUTE_PHASE_COLUMNS,
    ATTRIBUTE_PLACEHOLDER,
    ATTRIBUTE_RELATED_FIELDS,
    ATTRIBUTE_REQUIRED,
    ATTRIBUTE_RULE_AUTOFILL,
    ATTRIBUTE_RULE_AUTOFILL_READONLY,
    ATTRIBUTE_RULE_CONDITIONAL_VISIBILITY,
    ATTRIBUTE_RULE_UPDATE_AUTOFILL,
    ATTRIBUTE_SEARCHABLE,
    ATTRIBUTE_SUBGROUP,
    ATTRIBUTE_TYPE,
    ATTRIBUTE_UNIT,
    ATTRIBUTE_VALIDATION_REGEX,
    ATTRIBUTE_VIEW_PRIVILEGE,
    CARD_SECTION_NAME,
    CARD_SECTION_LOCATION,
    CARD_SECTION_DATE_FORMAT,
    CARD_SHOW_ON_MOBILE,
    CARD_EXTERNAL_DOCUMENT_FIELDS,
    CARD_EXTERNAL_DOCUMENT_SECTION,
    CARD_EXTERNAL_DOCUMENT_SECTION_INDEX,
    CALCULATIONS_COLUMN,
    CHOICE_OPTIONS_SHEET_NAME,
    CHOICES_SHEET_NAME,
    DEFAULT_SHEET_NAME,
    EXT_DATA_AD_KEY,
    EXT_DATA_AD_SOURCE,
    EXT_DATA_KEY_ATTRIBUTE,
    EXT_DATA_PARENT_KEY_ATTRIBUTE,
    EXT_DATA_SOURCE,
    EXT_DATA_SOURCE_KEY,
    HELP_IMG_LINK,
    HELP_LINK,
    HELP_TEXT,
    Phases,
    PROJECT_SIZE,
    PUBLIC_ATTRIBUTE,
    AttributeImporterException,
)

def get_mock_excel_rows(test_objects: list):
    """ Helper for generating mock Excel rows from a list of test objects.
    Args:
        test_objects (list): A list of dictionaries representing the test objects to generate mock Excel rows for.
            Each dictionary's keys are column names, and values are the corresponding cell values.
    Returns:
        list of lists: first list is header row, subsequent lists are data rows corresponding to the test_objects values.
    """
    if not test_objects:
        return []
    header = []
    for obj in test_objects:
        for key in obj.keys():
            if key not in header:
                header.append(key)
    rows = [header]
    for obj in test_objects:
        row = [obj.get(col, None) for col in header]
        rows.append(row)
    return rows

@pytest.mark.django_db
def test_project_phases_are_created(f_project_type, f_project_subtype):
    ai = AttributeImporter()
    ai.project_type = f_project_type
    ai.create_phases(f_project_subtype)

    assert CommonProjectPhase.objects.all().count() == 6

    for i in range(1, 7):
        assert CommonProjectPhase.objects.filter(index=i).exists()


def test_set_row_indexes():
    mock_rows = get_mock_excel_rows([
        {"name": "Attribute 1", "type": "date", "editable": "Kyllä"},
        {"name": "Attribute 2", "type": "text", "editable": "Kyllä"},
        {"name": "Attribute 3", "type": "number", "editable": "Ei"},
    ])

    ai = AttributeImporter()
    ai._set_row_indexes(mock_rows[0])
    assert ai.column_index is not None
    assert ai.column_index["name"] == 0
    assert ai.column_index["type"] == 1
    assert ai.column_index["editable"] == 2

def test_check_if_row_valid():
    mock_rows = get_mock_excel_rows([
        {ATTRIBUTE_NAME: "Attribute 1", ATTRIBUTE_TYPE: "date", "editable": "Kyllä"},
        {ATTRIBUTE_NAME: "Attribute 2", "editable": "Kyllä" },
        {ATTRIBUTE_NAME: None, ATTRIBUTE_TYPE: "number", "editable": "Ei"},
    ])
    ai = AttributeImporter()
    ai._set_row_indexes(mock_rows[0])
    assert ai._check_if_row_valid(mock_rows[1]) is True
    assert ai._check_if_row_valid(mock_rows[2]) is False
    assert ai._check_if_row_valid(mock_rows[3]) is False

def test_row_part_of_fieldset():
    from projects.importing.attribute import ATTRIBUTE_FIELDSET
    mock_rows = get_mock_excel_rows([
        {"name": "Attribute 1", "type": "date", ATTRIBUTE_FIELDSET: "Fieldset 1"},
        {"name": "Attribute 2", "type": "text"},
    ])

    ai = AttributeImporter()
    ai._set_row_indexes(mock_rows[0])
    assert ai._row_part_of_fieldset(mock_rows[1]) is True
    assert ai._row_part_of_fieldset(mock_rows[2]) is False

def test_get_attribute_row_identifier():
    mock_rows = get_mock_excel_rows([
        {ATTRIBUTE_IDENTIFIER: "single"},
        {ATTRIBUTE_IDENTIFIER: "multiple_word_identifier"},
        {ATTRIBUTE_IDENTIFIER: " whitespace "},
        {ATTRIBUTE_NAME: "Attribute 7"},
        {ATTRIBUTE_NAME: "Attribute-8"}
    ])

    ai = AttributeImporter()
    ai._set_row_indexes(mock_rows[0])
    assert ai._get_attribute_row_identifier(mock_rows[1]) == "single"
    assert ai._get_attribute_row_identifier(mock_rows[2]) == "multiple_word_identifier"
    assert ai._get_attribute_row_identifier(mock_rows[3]) == "whitespace"
    assert ai._get_attribute_row_identifier(mock_rows[4]) == "attribute_7"
    assert ai._get_attribute_row_identifier(mock_rows[5]) == "attribute_8"

def test_get_attribute_row_identifier_throws_on_invalid():
    mock_rows = get_mock_excel_rows([
        {ATTRIBUTE_IDENTIFIER: "with-dash"},
        {ATTRIBUTE_IDENTIFIER: "invalid identifier"},
        {ATTRIBUTE_IDENTIFIER: "v´+ery1nv<lid"},
    ])
    ai = AttributeImporter()
    ai._set_row_indexes(mock_rows[0])
    with pytest.raises(ValueError):
        ai._get_attribute_row_identifier(mock_rows[1])
    with pytest.raises(ValueError):
        ai._get_attribute_row_identifier(mock_rows[2])
    with pytest.raises(ValueError):
        ai._get_attribute_row_identifier(mock_rows[3])

def test_get_attribute_locations_parses_nested_location():
    phase_columns = ATTRIBUTE_PHASE_COLUMNS[Phases.START]
    rows = get_mock_excel_rows([
        {
            phase_columns[0]: "Section",
            phase_columns[1]: "Section ingress",
            phase_columns[2]: "1.2.3:4",
        }
    ])
    ai = AttributeImporter()
    ai._set_row_indexes(rows[0])

    assert ai._get_attribute_locations(rows[1], Phases.START.value) == {
        "label": "Section",
        "ingress": "Section ingress",
        "section_location": 10000,
        "field_location": 20000,
        "child_locations": [34000],
    }


def test_get_attribute_locations_without_child_location():
    phase_columns = ATTRIBUTE_PHASE_COLUMNS[Phases.OAS]
    rows = get_mock_excel_rows([
        {
            phase_columns[0]: "Section",
            phase_columns[1]: None,
            phase_columns[2]: "2.5",
        }
    ])
    ai = AttributeImporter()
    ai._set_row_indexes(rows[0])

    assert ai._get_attribute_locations(rows[1], Phases.OAS.value) == {
        "label": "Section",
        "ingress": None,
        "section_location": 20000,
        "field_location": 50000,
        "child_locations": [],
    }


@pytest.mark.parametrize("location", [None, "invalid"])
def test_get_attribute_locations_returns_none_for_invalid_location(location):
    phase_columns = ATTRIBUTE_PHASE_COLUMNS[Phases.REVISED_PROPOSAL]
    rows = get_mock_excel_rows([
        {
            phase_columns[0]: "Section",
            phase_columns[1]: None,
            phase_columns[2]: location,
        }
    ])
    ai = AttributeImporter()
    ai._set_row_indexes(rows[0])

    assert ai._get_attribute_locations(rows[1], Phases.REVISED_PROPOSAL.value) is None


def test_get_attribute_locations_returns_none_for_unknown_phase():
    ai = AttributeImporter()
    ai.column_index = {}

    assert ai._get_attribute_locations([], "Unknown phase") is None


@pytest.mark.parametrize(
    ("locations", "expected_index"),
    [
        ({"field_location": 20000, "child_locations": []}, 20000),
        ({"field_location": 20000, "child_locations": [34000]}, 23400),
        ({"field_location": 20000, "child_locations": [30000, 4000]}, 23040),
    ],
)
def test_calculate_index_preserves_field_and_nested_child_positions(
    locations, expected_index
):
    assert AttributeImporter.calculate_index(locations) == expected_index


@pytest.mark.parametrize(
    ("locations", "expected_index"),
    [
        ([], 0),
        (["1"], 1000),
        (["1", "2", "3"], 1203),
        (["2", "5", "1", "7"], 2508),
    ],
)
def test_calculate_deadline_index_scales_each_position_by_depth(
    locations, expected_index
):
    assert AttributeImporter.calculate_deadline_index(locations) == expected_index


@pytest.mark.django_db
def test_phase_sections_rebuild_and_link_only_attributes_for_matching_subtype(
    f_project_subtype, f_project_phase_1
):
    section_columns = ATTRIBUTE_PHASE_COLUMNS[Phases.START]
    existing_section = ProjectPhaseSection.objects.create(
        phase=f_project_phase_1,
        name="Obsolete section",
        index=90000,
    )
    attributes = {
        identifier: Attribute.objects.create(
            name=identifier.replace("_", " "),
            identifier=identifier,
            value_type=Attribute.TYPE_SHORT_STRING,
        )
        for identifier in ["phase_field_m_first", "phase_field_xl", "phase_field_m_second"]
    }
    rows = get_mock_excel_rows(
        [
            {
                ATTRIBUTE_IDENTIFIER: "phase_field_m_first",
                PROJECT_SIZE: "M",
                section_columns[0]: "Project overview",
                section_columns[1]: "Intro text",
                section_columns[2]: "1.2.3.4",
            },
            {
                ATTRIBUTE_IDENTIFIER: "phase_field_xl",
                PROJECT_SIZE: "XL",
                section_columns[0]: "Project overview",
                section_columns[1]: "Intro text",
                section_columns[2]: "1.2.3.5",
            },
            {
                ATTRIBUTE_IDENTIFIER: "phase_field_m_second",
                PROJECT_SIZE: "M",
                section_columns[0]: "Project overview",
                section_columns[1]: "Intro text",
                section_columns[2]: "1.2.3.6",
            },
        ]
    )
    importer = AttributeImporter()
    importer._set_row_indexes(rows[0])

    importer._create_sections(rows[1:], f_project_subtype)
    importer._create_attribute_section_links(rows[1:], f_project_subtype)

    sections = list(f_project_phase_1.sections.all())
    assert len(sections) == 1
    section = sections[0]
    assert (section.name, section.ingress, section.index) == (
        "Project overview",
        "Intro text",
        10000,
    )
    assert not ProjectPhaseSection.objects.filter(pk=existing_section.pk).exists()

    links = list(
        ProjectPhaseSectionAttribute.objects.filter(section=section)
        .order_by("index")
        .values_list("attribute__identifier", "index")
    )
    assert links == [
        ("phase_field_m_first", 23400),
        ("phase_field_m_second", 23600),
    ]
    assert "phase_field_xl" not in [attribute_id for attribute_id, _ in links]


@pytest.mark.django_db
def test_deadline_sections_assign_owner_admin_roles_and_filter_other_subtypes(
    f_project_type, f_project_subtype, f_project_phase_1
):
    f_project_subtype.name = "M"
    f_project_subtype.save(update_fields=["name"])
    other_subtype = ProjectSubtype.objects.create(
        name="XS",
        project_type=f_project_type,
        index=1,
    )
    other_phase = ProjectPhase.objects.create(
        common_project_phase=f_project_phase_1.common_project_phase,
        project_subtype=other_subtype,
        index=0,
    )
    admin_attribute = Attribute.objects.create(
        name="Admin date",
        identifier="deadline_admin_date",
        value_type=Attribute.TYPE_DATE,
    )
    owner_attribute = Attribute.objects.create(
        name="Owner date",
        identifier="deadline_owner_date",
        value_type=Attribute.TYPE_DATE,
    )
    excluded_attribute = Attribute.objects.create(
        name="Other subtype date",
        identifier="deadline_other_subtype_date",
        value_type=Attribute.TYPE_DATE,
    )
    Deadline.objects.create(
        abbreviation="ADM",
        attribute=admin_attribute,
        phase=f_project_phase_1,
        subtype=f_project_subtype,
    )
    Deadline.objects.create(
        abbreviation="OWN",
        attribute=owner_attribute,
        phase=f_project_phase_1,
        subtype=f_project_subtype,
    )
    Deadline.objects.create(
        abbreviation="XS",
        attribute=excluded_attribute,
        phase=other_phase,
        subtype=other_subtype,
    )
    columns = [
        ATTRIBUTE_IDENTIFIER,
        PROJECT_SIZE,
        ATTRIBUTE_DEADLINE_SECTION_COLUMNS["admin"],
        ATTRIBUTE_DEADLINE_SECTION_COLUMNS["owner"],
    ]
    rows = get_mock_excel_rows(
        [
            {
                ATTRIBUTE_IDENTIFIER: "deadline_admin_date",
                PROJECT_SIZE: "M",
                ATTRIBUTE_DEADLINE_SECTION_COLUMNS["admin"]: "Käynnistys; 1.2.3",
            },
            {
                ATTRIBUTE_IDENTIFIER: "deadline_owner_date",
                PROJECT_SIZE: "M",
                ATTRIBUTE_DEADLINE_SECTION_COLUMNS["owner"]: "Käynnistys; 2.3.4",
            },
            {
                ATTRIBUTE_IDENTIFIER: "deadline_other_subtype_date",
                PROJECT_SIZE: "XS",
                ATTRIBUTE_DEADLINE_SECTION_COLUMNS["admin"]: "Käynnistys; 3.4.5",
            },
        ]
    )
    importer = AttributeImporter()
    importer._set_row_indexes(columns)

    importer._create_deadline_sections(rows[1:], f_project_subtype)

    section = ProjectPhaseDeadlineSection.objects.get(phase=f_project_phase_1)
    assignments = {
        item.attribute.identifier: (item.admin_field, item.owner_field, item.index)
        for item in ProjectPhaseDeadlineSectionAttribute.objects.filter(section=section)
    }
    assert assignments == {
        "deadline_admin_date": (True, False, 1203),
        "deadline_owner_date": (False, True, 2304),
    }
    assert not ProjectPhaseDeadlineSection.objects.filter(phase=other_phase).exists()


def make_choice_importer(choice_rows, choices_ref):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = CHOICES_SHEET_NAME
    for choice_row in choice_rows:
        worksheet.append(choice_row)

    importer = AttributeImporter()
    importer.workbook = workbook
    rows = get_mock_excel_rows([{ATTRIBUTE_CHOICES_REF: choices_ref}])
    importer._set_row_indexes(rows[0])
    return importer, rows[1]


@pytest.mark.django_db
def test_attribute_choices_use_referenced_column():
    attribute = Attribute.objects.create(
        name="Status",
        identifier="choice_status",
        value_type=Attribute.TYPE_CHOICE,
    )
    importer, row = make_choice_importer(
        [
            ["Other list", "Status"],
            ["other_value", "draft"],
            ["other_approved", "approved"],
            ["ignored", None],
            ["must_not_be_imported", None],
        ],
        "Status",
    )

    created_count = importer._create_attribute_choices(
        attribute,
        row,
        {"draft": "Draft", "approved": "Approved"},
    )

    choices = list(
        attribute.value_choices.order_by("index").values_list(
            "identifier", "value", "index"
        )
    )
    assert created_count == 2
    assert choices == [("draft", "Draft", 1), ("approved", "Approved", 2)]


@pytest.mark.django_db
def test_attribute_choices_accept_slash_separated_identifiers_without_column():
    attribute = Attribute.objects.create(
        name="Boolean-like choice",
        identifier="choice_boolean_like",
        value_type=Attribute.TYPE_CHOICE,
    )
    importer, row = make_choice_importer([["Unused column"], ["unused"]], "yes/no")

    created_count = importer._create_attribute_choices(attribute, row, {})

    choices = list(
        attribute.value_choices.order_by("index").values_list(
            "identifier", "value", "index"
        )
    )
    assert created_count == 2
    assert choices == [("yes", "yes", 0), ("no", "no", 1)]


@pytest.mark.django_db
def test_attribute_choice_identifier_change_reuses_matching_legacy_choice():
    attribute = Attribute.objects.create(
        name="Status",
        identifier="choice_status_legacy",
        value_type=Attribute.TYPE_CHOICE,
    )
    importer, row = make_choice_importer(
        [["Status"], ["new_draft"]],
        "Status",
    )
    display_value = "Draft"
    legacy_identifier = importer._get_identifier_for_value(display_value)
    existing_choice = AttributeValueChoice.objects.create(
        attribute=attribute,
        identifier="old_draft",
        legacy_identifier=legacy_identifier,
        value=display_value,
        index=0,
    )

    created_count = importer._create_attribute_choices(
        attribute,
        row,
        {"new_draft": display_value},
    )

    existing_choice.refresh_from_db()
    assert created_count == 0
    assert attribute.value_choices.count() == 1
    assert (
        existing_choice.identifier,
        existing_choice.value,
        existing_choice.index,
        existing_choice.legacy_identifier,
    ) == ("new_draft", display_value, 1, legacy_identifier)


# Every column `_create_attributes` reads from a row, in header order.
ATTRIBUTE_ROW_COLUMNS = [
    ATTRIBUTE_NAME,
    ATTRIBUTE_IDENTIFIER,
    ATTRIBUTE_TYPE,
    ATTRIBUTE_CHOICES_REF,
    ATTRIBUTE_UNIT,
    ATTRIBUTE_BROADCAST_CHANGES,
    ATTRIBUTE_REQUIRED,
    ATTRIBUTE_DATA_RETENTION,
    ATTRIBUTE_MULTIPLE_CHOICE,
    ATTRIBUTE_SEARCHABLE,
    ATTRIBUTE_RELATED_FIELDS,
    ATTRIBUTE_LINKED_FIELDS,
    ATTRIBUTE_RULE_CONDITIONAL_VISIBILITY,
    ATTRIBUTE_RULE_AUTOFILL,
    ATTRIBUTE_RULE_AUTOFILL_READONLY,
    ATTRIBUTE_RULE_UPDATE_AUTOFILL,
    ATTRIBUTE_CHARACTER_LIMIT,
    ATTRIBUTE_VALIDATION_REGEX,
    ATTRIBUTE_HIGHLIGHT_GROUP,
    ATTRIBUTE_EDIT_PRIVILEGE,
    ATTRIBUTE_VIEW_PRIVILEGE,
    ATTRIBUTE_ERROR,
    ATTRIBUTE_PLACEHOLDER,
    ATTRIBUTE_ASSISTIVE_TEXT,
    ATTRIBUTE_ERROR_TEXT,
    ATTRIBUTE_FIELD_ROLE,
    ATTRIBUTE_FIELD_SUBROLE,
    ATTRIBUTE_FIELDSET_TOTAL,
    ATTRIBUTE_GROUP,
    ATTRIBUTE_SUBGROUP,
    ATTRIBUTE_API_VISIBILITY,
    PUBLIC_ATTRIBUTE,
    HELP_TEXT,
    HELP_LINK,
    HELP_IMG_LINK,
    EXT_DATA_SOURCE,
    EXT_DATA_SOURCE_KEY,
    EXT_DATA_PARENT_KEY_ATTRIBUTE,
    EXT_DATA_AD_KEY,
    CALCULATIONS_COLUMN,
]


def build_attribute_row(overrides):
    """Build one spreadsheet-style attribute row, defaulting every column
    `_create_attributes` reads so each test only has to override what it cares about."""
    values = {column: None for column in ATTRIBUTE_ROW_COLUMNS}
    values[ATTRIBUTE_IDENTIFIER] = ""
    values.update(overrides)
    return [values[column] for column in ATTRIBUTE_ROW_COLUMNS]


def make_attribute_importer():
    workbook = Workbook()
    choice_options_sheet = workbook.active
    choice_options_sheet.title = CHOICE_OPTIONS_SHEET_NAME
    # _get_values_by_identifier reads this sheet on every import, even when no row
    # in the import references a choice list.
    choice_options_sheet.append(["#", "unused_identifier", "unused_value"])

    importer = AttributeImporter()
    importer.workbook = workbook
    importer._set_row_indexes(ATTRIBUTE_ROW_COLUMNS)
    return importer


@pytest.mark.django_db
def test_create_attributes_creates_updates_and_deletes_stale_attribute_in_one_import():
    existing_attribute = Attribute.objects.create(
        name="Old name",
        identifier="existing_attr",
        value_type=Attribute.TYPE_SHORT_STRING,
        required=False,
        public=False,
    )
    AttributeValueChoice.objects.create(
        attribute=existing_attribute,
        identifier="existing_choice",
        value="Existing choice",
        index=0,
    )
    stale_attribute = Attribute.objects.create(
        name="Stale attribute",
        identifier="stale_attr",
        value_type=Attribute.TYPE_SHORT_STRING,
    )

    rows = [
        build_attribute_row({
            ATTRIBUTE_NAME: "Updated name",
            ATTRIBUTE_IDENTIFIER: "existing_attr",
            ATTRIBUTE_TYPE: "Kokonaisluvun syöttö.",
            ATTRIBUTE_REQUIRED: "kyllä",
            PUBLIC_ATTRIBUTE: "kyllä",
        }),
        build_attribute_row({
            ATTRIBUTE_NAME: "Brand new attribute",
            ATTRIBUTE_IDENTIFIER: "new_attr",
            ATTRIBUTE_TYPE: "Kokonaisluvun syöttö.",
        }),
    ]

    importer = make_attribute_importer()
    result = importer._create_attributes(rows)

    assert result["created"] == 1
    assert result["updated"] == 1
    assert result["deleted"] == 1

    existing_attribute.refresh_from_db()
    assert existing_attribute.name == "Updated name"
    assert existing_attribute.value_type == Attribute.TYPE_INTEGER
    assert existing_attribute.required is True
    assert existing_attribute.public is True
    # Row had no choices reference, so _create_attributes must clear old choices.
    assert existing_attribute.value_choices.count() == 0

    new_attribute = Attribute.objects.get(identifier="new_attr")
    assert new_attribute.name == "Brand new attribute"
    assert new_attribute.value_type == Attribute.TYPE_INTEGER

    assert not Attribute.objects.filter(identifier="stale_attr").exists()


def build_key_relation_row(overrides):
    """Build one row with only the columns `_create_attribute_key_relations` reads."""
    columns = [
        ATTRIBUTE_IDENTIFIER,
        EXT_DATA_KEY_ATTRIBUTE,
        EXT_DATA_AD_SOURCE,
        ATTRIBUTE_AUTO_VALUE_KEY_FIELD,
        ATTRIBUTE_AUTO_VALUE_MAPPING,
    ]
    values = {column: None for column in columns}
    values[ATTRIBUTE_IDENTIFIER] = ""
    values.update(overrides)
    return columns, [values[column] for column in columns]


@pytest.mark.django_db
def test_create_attribute_key_relations_wires_and_translates_auto_value_mapping():
    key_source_attr = Attribute.objects.create(name="Key source", identifier="key_source_attr")
    ad_source_attr = Attribute.objects.create(name="AD source", identifier="ad_source_attr")
    selected_role = Attribute.objects.create(
        name="Selected role", identifier="selected_role", value_type=Attribute.TYPE_CHOICE
    )
    AttributeValueChoice.objects.create(attribute=selected_role, identifier="role_a", value="Role A", index=0)
    AttributeValueChoice.objects.create(attribute=selected_role, identifier="role_b", value="Role B", index=1)
    contact_person = Attribute.objects.create(name="Contact person", identifier="contact_person")

    # Stale relation that must be cleared since it isn't part of this import.
    stale_attr = Attribute.objects.create(
        name="Stale", identifier="stale_attr_with_old_key", key_attribute=key_source_attr
    )
    # Stale auto-value mapping that must be removed since it isn't referenced by this import.
    stale_target = Attribute.objects.create(name="Stale target", identifier="stale_target_attr")
    stale_auto_attr = AttributeAutoValue.objects.create(
        value_attribute=stale_target, key_attribute=selected_role
    )
    stale_mapping = AttributeAutoValueMapping.objects.create(
        auto_attr=stale_auto_attr, key_str="role_a", value_str="Old contact"
    )

    columns, row = build_key_relation_row({
        ATTRIBUTE_IDENTIFIER: "contact_person",
        EXT_DATA_KEY_ATTRIBUTE: "key_source_attr",
        EXT_DATA_AD_SOURCE: "ad_source_attr",
        ATTRIBUTE_AUTO_VALUE_KEY_FIELD: "selected_role",
        ATTRIBUTE_AUTO_VALUE_MAPPING: '"Role A": "Contact A"; "Role B": "Contact B"',
    })
    ai = AttributeImporter()
    ai._set_row_indexes(columns)
    ai._create_attribute_key_relations([row])

    contact_person.refresh_from_db()
    assert contact_person.key_attribute_id == key_source_attr.id
    assert contact_person.ad_key_attribute_id == ad_source_attr.id

    auto_attr = AttributeAutoValue.objects.get(value_attribute=contact_person)
    assert auto_attr.key_attribute_id == selected_role.id
    # Mapping keys use value_choice identifiers, translated from the sheet's display text.
    mappings = {m.key_str: m.value_str for m in auto_attr.value_map.all()}
    assert mappings == {"role_a": "Contact A", "role_b": "Contact B"}

    stale_attr.refresh_from_db()
    assert stale_attr.key_attribute is None
    assert not AttributeAutoValue.objects.filter(pk=stale_auto_attr.pk).exists()
    assert not AttributeAutoValueMapping.objects.filter(pk=stale_mapping.pk).exists()


@pytest.mark.django_db
def test_create_subtypes_orders_by_size_and_dedupes_case_insensitively(f_project_type):
    rows = get_mock_excel_rows([
        {PROJECT_SIZE: "XL"},
        {PROJECT_SIZE: "M, XS"},
        {PROJECT_SIZE: "m"},  # duplicate of "M" in different case, must not create a second subtype
        {PROJECT_SIZE: "kaikki"},  # applies to every subtype, must not create a subtype of its own
    ])
    ai = AttributeImporter()
    ai.project_type = f_project_type
    ai._set_row_indexes(rows[0])

    subtypes = ai.create_subtypes(rows[1:])

    assert [s.name for s in subtypes] == ["XS", "M", "XL"]
    assert [s.index for s in subtypes] == [0, 1, 2]
    assert ProjectSubtype.objects.filter(project_type=f_project_type).count() == 3


@pytest.mark.parametrize(
    ("cell_content", "expected_subtypes"),
    [
        (None, ["kaikki"]),
        ("", ["kaikki"]),
        ("kaikki", ["kaikki"]),
        ("XL", ["xl"]),
        ("S, M, L", ["s", "m", "l"]),
        ("XS,XL", ["xs", "xl"]),
    ],
)
def test_get_subtypes_from_cell_parses_project_size_lists(cell_content, expected_subtypes):
    ai = AttributeImporter()
    assert ai.get_subtypes_from_cell(cell_content) == expected_subtypes


@pytest.mark.parametrize(
    ("calculations_string", "expected"),
    [
        (None, (False, None)),
        ("ei", (False, None)),
        ("kerrosala", (True, ["kerrosala"])),
        (
            "kerrosala + lisakerrosala - varattu_kerrosala",
            (True, ["kerrosala", "+", "lisakerrosala", "-", "varattu_kerrosala"]),
        ),
    ],
)
def test_get_generated_calculations_tokenizes_formula_string(calculations_string, expected):
    ai = AttributeImporter()
    ai._set_row_indexes([CALCULATIONS_COLUMN])
    assert ai._get_generated_calculations([calculations_string]) == expected


def test_extract_data_from_workbook_raises_for_missing_sheet():
    workbook = Workbook()
    ai = AttributeImporter(options={"sheet": "Nonexistent sheet"})

    with pytest.raises(AttributeImporterException):
        ai._extract_data_from_workbook(workbook)


def test_extract_data_from_workbook_raises_for_unexpected_a1_value():
    workbook = Workbook()
    workbook.active.title = DEFAULT_SHEET_NAME
    workbook.active.append(["Not the expected header"])
    ai = AttributeImporter(options={})

    with pytest.raises(AttributeImporterException):
        ai._extract_data_from_workbook(workbook)


def test_extract_data_from_workbook_returns_rows_after_header():
    workbook = Workbook()
    workbook.active.title = DEFAULT_SHEET_NAME
    workbook.active.append([ATTRIBUTE_NAME, ATTRIBUTE_TYPE])
    workbook.active.append(["Attribute 1", "date"])
    workbook.active.append(["Attribute 2", "text"])
    ai = AttributeImporter(options={})

    rows = ai._extract_data_from_workbook(workbook)

    assert rows == [["Attribute 1", "date"], ["Attribute 2", "text"]]


def test_open_workbook_raises_importer_exception_for_missing_file():
    ai = AttributeImporter()

    with pytest.raises(AttributeImporterException):
        ai._open_workbook("/nonexistent/path/does-not-exist.xlsx")


@pytest.mark.django_db
@pytest.mark.parametrize("value_type", [Attribute.TYPE_INTEGER, Attribute.TYPE_DECIMAL])
def test_generated_calculation_accepts_existing_numeric_inputs(value_type):
    Attribute.objects.create(
        name="Area",
        identifier="calculation_area",
        value_type=value_type,
    )
    Attribute.objects.create(
        name="Units",
        identifier="calculation_units",
        value_type=value_type,
    )
    Attribute.objects.create(
        name="Calculated total",
        identifier="calculated_total",
        value_type=value_type,
        generated=True,
        calculations=["calculation_area", "+", "calculation_units"],
    )
    # Does not throw an exception because all input attributes exist and are numeric
    AttributeImporter()._validate_generated_attributes()


@pytest.mark.django_db
def test_generated_calculation_rejects_missing_input_attribute():
    Attribute.objects.create(
        name="Calculated total",
        identifier="calculated_total",
        value_type=Attribute.TYPE_INTEGER,
        generated=True,
        calculations=["missing_input", "+", "also_missing"],
    )
    ai = AttributeImporter()
    with pytest.raises(Exception, match="Could not add attribute calculated_total"):
        ai._validate_generated_attributes()


@pytest.mark.django_db
def test_generated_calculation_rejects_non_numeric_input_attribute():
    Attribute.objects.create(
        name="Description",
        identifier="calculation_description",
        value_type=Attribute.TYPE_SHORT_STRING,
    )
    Attribute.objects.create(
        name="Calculated total",
        identifier="calculated_total",
        value_type=Attribute.TYPE_INTEGER,
        generated=True,
        calculations=["calculation_description"],
    )

    ai = AttributeImporter()
    with pytest.raises(Exception, match="Could not add attribute calculated_total"):
        ai._validate_generated_attributes()


@pytest.mark.django_db
def test_non_generated_calculation_is_not_validated_by_importer():
    Attribute.objects.create(
        name="Calculated total",
        identifier="not_generated_total",
        value_type=Attribute.TYPE_INTEGER,
        generated=False,
        calculations=["missing_input"],
    )

    AttributeImporter()._validate_generated_attributes()

def test_parse_condition():
    ai = AttributeImporter()
    assert ai._parse_condition("oas_mielipiteet_maara > 0") == {
        "variable": "oas_mielipiteet_maara",
        "operator": ">",
        "comparison_value": "0",
        "comparison_value_type": "number",
    }
    assert ai._parse_condition("jarjestetaan_luonnosvaiheessa_tilaisuus_3") == {
        "variable": "jarjestetaan_luonnosvaiheessa_tilaisuus_3",
        "operator": "==",
        "comparison_value": True,
        "comparison_value_type": "boolean",
    }
    assert ai._parse_condition("!maanomistus_kaupunki") == {
        "variable": "maanomistus_kaupunki",
        "operator": "!=",
        "comparison_value": True,
        "comparison_value_type": "boolean",
    }
    assert ai._parse_condition("kaavaprosessin_kokoluokka in [\"L\", \"XL\"]") == {
        "variable": "kaavaprosessin_kokoluokka",
        "operator": "in",
        "comparison_value": '["L", "XL"]',
        "comparison_value_type": "list<string>",
    }
    assert ai._parse_condition("kaavaprosessin_kokoluokka not in [\"S\", \"M\"]") == {
        "variable": "kaavaprosessin_kokoluokka",
        "operator": "not in",
        "comparison_value": '["S", "M"]',
        "comparison_value_type": "list<string>",
    }
    assert ai._parse_condition("lautakunta_paatti_ehdotus == \"paatos_kaavaehdotuksesta_asia_jai_poydalle\"") == {
        "variable": "lautakunta_paatti_ehdotus",
        "operator": "==",
        "comparison_value": "paatos_kaavaehdotuksesta_asia_jai_poydalle",
        "comparison_value_type": "string",
    }

def test_parse_autofill_rule():
    ai = AttributeImporter()

    # Test rule for autofilling a value with a simple "ei" condition
    assert ai._parse_autofill_rule("ei", Attribute.TYPE_SHORT_STRING) == None

    # Test rule for autofilling a hardcoded value
    assert ai._parse_autofill_rule("{% if !maanomistus_kaupunki %} 0 {% endif %}", Attribute.TYPE_INTEGER) == [
        {
            "variables": [],
            "conditions": [
                {
                    "variable": "maanomistus_kaupunki",
                    "operator": "!=",
                    "comparison_value": True,
                    "comparison_value_type": "boolean",
                }
            ],
            "then_branch": "0",
            "else_branch": None
        }
    ]
    # Test rule for autofilling a value based on a variable
    assert ai._parse_autofill_rule("{{milta_ulkopuolisilta_pyydetaan_lausunto}} {% if milta_ulkopuolisilta_pyydetaan_lausunto %} {% endif %}", Attribute.TYPE_CHOICE) == [
        {
            "variables": ["milta_ulkopuolisilta_pyydetaan_lausunto"],
            "conditions": [
                {
                    "variable": "milta_ulkopuolisilta_pyydetaan_lausunto",
                    "operator": "==",
                    "comparison_value": True,
                    "comparison_value_type": "boolean",
                }
            ],
            "then_branch": "",
            "else_branch": None
        }
    ]
    # Test rule for autofilling a value based on multiple conditions
    assert ai._parse_autofill_rule("{% if periaatteet_mielipiteet_maara > 0 or oas_mielipiteet_maara > 0 %} kyllä {% endif %}",
        Attribute.TYPE_CHOICE
    ) == [
        {
            "variables": [],
            "conditions": [
                {
                    "variable": "periaatteet_mielipiteet_maara",
                    "operator": ">",
                    "comparison_value": '0',
                    "comparison_value_type": "number",
                },
            ],
            "then_branch": "kylla",
            "else_branch": None
        },
        {
            "variables": [],
            "conditions": [
                {
                    "variable": "oas_mielipiteet_maara",
                    "operator": ">",
                    "comparison_value": '0',
                    "comparison_value_type": "number",
                },
            ],
            "then_branch": "kylla",
            "else_branch": None
        },
    ]
    # Test boolean conversion kylla/ei to True/False
    assert ai._parse_autofill_rule("{% if condition %} kyllä {% endif %}", Attribute.TYPE_BOOLEAN) == [
        {
            "variables": [],
            "conditions": [
                {
                    "variable": "condition",
                    "operator": "==",
                    "comparison_value": True,
                    "comparison_value_type": "boolean",
                }
            ],
            "then_branch": True,
            "else_branch": None
        }
    ]
    assert ai._parse_autofill_rule("{% if !condition %} ei {% endif %}", Attribute.TYPE_BOOLEAN) == [
        {
            "variables": [],
            "conditions": [
                {
                    "variable": "condition",
                    "operator": "!=",
                    "comparison_value": True,
                    "comparison_value_type": "boolean",
                }
            ],
            "then_branch": False,
            "else_branch": None
        }
    ]

def test_parse_autofill_readonly():
    # Tests parsing rules for "is autofilled column editable"- column
    ai = AttributeImporter()
    assert ai._parse_autofill_readonly("ei") == True
    assert ai._parse_autofill_readonly("Automaattiseti muodostunutta tietoa ei voi muokata") == True
    assert ai._parse_autofill_readonly("kyllä") == False
    assert ai._parse_autofill_readonly("") == False


@pytest.mark.django_db
def test_document_link_import_replaces_old_links_and_skips_unconfigured_rows():
    fieldset = Attribute.objects.create(name="Documents", identifier="documents")
    document_name = Attribute.objects.create(name="Name", identifier="document_name")
    document_link = Attribute.objects.create(name="Link", identifier="document_link")
    old_fieldset = Attribute.objects.create(name="Old documents", identifier="old_documents")
    old_section = DocumentLinkSection.objects.create(name="Old", index=0)
    DocumentLinkFieldSet.objects.create(
        section=old_section,
        fieldset_attribute=old_fieldset,
        document_name_attribute=document_name,
        document_link_attribute=document_link,
    )

    rows = get_mock_excel_rows([
        {
            ATTRIBUTE_IDENTIFIER: "documents",
            CARD_EXTERNAL_DOCUMENT_SECTION: "Planning documents",
            CARD_EXTERNAL_DOCUMENT_SECTION_INDEX: 3,
            CARD_EXTERNAL_DOCUMENT_FIELDS: "document_name;;document_link",
        },
        {
            ATTRIBUTE_IDENTIFIER: "old_documents",
            CARD_EXTERNAL_DOCUMENT_SECTION: None,
            CARD_EXTERNAL_DOCUMENT_SECTION_INDEX: None,
            CARD_EXTERNAL_DOCUMENT_FIELDS: None,
        },
    ])
    importer = AttributeImporter()
    importer._set_row_indexes(rows[0])

    importer._create_document_link_sections(rows[1:])

    link_set = DocumentLinkFieldSet.objects.get()
    assert DocumentLinkSection.objects.get(pk=link_set.section_id).name == "Planning documents"
    assert link_set.fieldset_attribute == fieldset
    assert link_set.document_name_attribute == document_name
    assert link_set.document_custom_name_attribute is None
    assert link_set.document_link_attribute == document_link


@pytest.mark.django_db
def test_card_section_import_preserves_order_and_visibility_and_skips_empty_rows():
    attribute = Attribute.objects.create(name="Project name", identifier="card_project_name")
    obsolete_attribute = Attribute.objects.create(name="Obsolete", identifier="card_obsolete")
    obsolete_section = ProjectCardSection.objects.create(name="Obsolete section", index=9)
    ProjectCardSectionAttribute.objects.create(
        attribute=obsolete_attribute,
        section=obsolete_section,
    )
    rows = get_mock_excel_rows([
        {
            ATTRIBUTE_IDENTIFIER: "card_project_name",
            CARD_SECTION_NAME: "Perustiedot",
            CARD_SECTION_LOCATION: "2.4",
            CARD_SECTION_DATE_FORMAT: "updated",
            CARD_SHOW_ON_MOBILE: "ei",
        },
        {
            ATTRIBUTE_IDENTIFIER: "card_obsolete",
            CARD_SECTION_NAME: "ei",
            CARD_SECTION_LOCATION: None,
            CARD_SECTION_DATE_FORMAT: None,
            CARD_SHOW_ON_MOBILE: None,
        },
    ])
    importer = AttributeImporter()
    importer._set_row_indexes(rows[0])

    importer._create_card_sections(rows[1:])

    section = ProjectCardSection.objects.get()
    section_attribute = ProjectCardSectionAttribute.objects.get()
    assert (section.name, section.index, section.key) == ("Perustiedot", 2, "perustiedot")
    assert section_attribute.attribute == attribute
    assert (section_attribute.index, section_attribute.date_format, section_attribute.show_on_mobile) == (
        4,
        "updated",
        False,
    )


@pytest.mark.django_db
def test_fieldset_import_links_nested_children_and_keeps_unindexed_children_unindexed(
    f_project_subtype, f_project_phase_1
):
    section_columns = ATTRIBUTE_PHASE_COLUMNS[Phases.START]
    fieldset = Attribute.objects.create(name="Contacts", identifier="contact_fieldset")
    nested_child = Attribute.objects.create(name="Nested contact", identifier="nested_contact")
    direct_child = Attribute.objects.create(name="Direct contact", identifier="direct_contact")
    rows = get_mock_excel_rows([
        {
            ATTRIBUTE_IDENTIFIER: "nested_contact",
            ATTRIBUTE_FIELDSET: "contact_fieldset",
            section_columns[0]: "Contacts",
            section_columns[1]: None,
            section_columns[2]: "1.2.3:4",
        },
        {
            ATTRIBUTE_IDENTIFIER: "direct_contact",
            ATTRIBUTE_FIELDSET: "contact_fieldset",
            section_columns[0]: "Contacts",
            section_columns[1]: None,
            section_columns[2]: "1.2",
        },
    ])
    importer = AttributeImporter()
    importer._set_row_indexes(rows[0])

    importer._create_fieldset_links(f_project_subtype, rows[1:])

    links = {
        link.attribute_target_id: link
        for link in FieldSetAttribute.objects.filter(attribute_source=fieldset)
    }
    assert set(links) == {nested_child.pk, direct_child.pk}
    assert ProjectPhaseFieldSetAttributeIndex.objects.filter(
        attribute=links[nested_child.pk],
        phase=f_project_phase_1,
        index=34000,
    ).exists()
    assert not ProjectPhaseFieldSetAttributeIndex.objects.filter(
        attribute=links[direct_child.pk],
    ).exists()


@pytest.mark.django_db
def test_floor_area_import_filters_subtypes_and_builds_matrix_cells(f_project_subtype):
    f_project_subtype.name = "M"
    f_project_subtype.save(update_fields=["name"])
    section = ProjectFloorAreaSection.objects.create(
        project_subtype=f_project_subtype,
        name="Building area",
        index=1,
    )
    included = Attribute.objects.create(name="Area matrix", identifier="included_area_matrix")
    other_subtype = Attribute.objects.create(name="XL area", identifier="xl_area")
    excluded = Attribute.objects.create(name="Excluded area", identifier="excluded_area")
    rows = get_mock_excel_rows([
        {
            ATTRIBUTE_IDENTIFIER: "included_area_matrix",
            PROJECT_SIZE: "M",
            ATTRIBUTE_FLOOR_AREA_SECTION: "Building area",
            ATTRIBUTE_FLOOR_AREA_SECTION_MATRIX_ROW: "Residential\nCommercial",
            ATTRIBUTE_FLOOR_AREA_SECTION_MATRIX_CELL: "New\nExisting",
        },
        {
            ATTRIBUTE_IDENTIFIER: "xl_area",
            PROJECT_SIZE: "XL",
            ATTRIBUTE_FLOOR_AREA_SECTION: "Building area",
            ATTRIBUTE_FLOOR_AREA_SECTION_MATRIX_ROW: "ei",
            ATTRIBUTE_FLOOR_AREA_SECTION_MATRIX_CELL: "ei",
        },
        {
            ATTRIBUTE_IDENTIFIER: "excluded_area",
            PROJECT_SIZE: "M",
            ATTRIBUTE_FLOOR_AREA_SECTION: "ei",
            ATTRIBUTE_FLOOR_AREA_SECTION_MATRIX_ROW: "ei",
            ATTRIBUTE_FLOOR_AREA_SECTION_MATRIX_CELL: "ei",
        },
    ])
    importer = AttributeImporter()
    importer._set_row_indexes(rows[0])

    importer._create_floor_area_attribute_section_links(rows[1:], f_project_subtype)

    section_attributes = list(
        ProjectFloorAreaSectionAttribute.objects.filter(section=section).order_by("index")
    )
    assert [(item.attribute, item.index) for item in section_attributes] == [(included, 0)]
    structure = ProjectFloorAreaSectionAttributeMatrixStructure.objects.get(section=section)
    assert structure.row_names == ["Residential", "Commercial"]
    assert structure.column_names == ["New", "Existing"]
    cells = set(
        ProjectFloorAreaSectionAttributeMatrixCell.objects.filter(
            attribute=section_attributes[0]
        ).values_list("row", "column")
    )
    assert cells == {(0, 0), (0, 1), (1, 0), (1, 1)}
    assert other_subtype not in [item.attribute for item in section_attributes]
    assert excluded not in [item.attribute for item in section_attributes]


@pytest.mark.django_db
def test_attribute_categorization_import_updates_values_and_removes_stale_rows():
    attribute = Attribute.objects.create(name="Project scale", identifier="project_scale")
    phase_by_name = {
        phase.value: CommonProjectPhase.objects.create(name=phase.value)
        for phase in Phases
    }
    start_phase = phase_by_name[Phases.START.value]
    existing = AttributeCategorization.objects.create(
        attribute=attribute,
        common_project_phase=start_phase,
        includes_principles=False,
        includes_draft=False,
        value="Old value",
    )
    stale = AttributeCategorization.objects.create(
        attribute=attribute,
        common_project_phase=phase_by_name[Phases.OAS.value],
        includes_principles=False,
        includes_draft=False,
        value="Stale value",
    )
    start_column = ATTRIBUTE_CATEGORIZATION_COLUMNS[0][0]
    columns = list(dict.fromkeys(
        [ATTRIBUTE_IDENTIFIER]
        + [column for column, _, _ in ATTRIBUTE_CATEGORIZATION_COLUMNS]
    ))
    row_values = {column: None for column in columns}
    row_values[ATTRIBUTE_IDENTIFIER] = "project_scale"
    row_values[start_column] = "Scale category"
    rows = get_mock_excel_rows([row_values])
    importer = AttributeImporter()
    importer._set_row_indexes(rows[0])

    importer._create_attribute_categorizations(rows[1:])

    existing.refresh_from_db()
    assert existing.value == "Scale category"
    start_categorizations = AttributeCategorization.objects.filter(
        attribute=attribute,
        common_project_phase=start_phase,
    )
    assert set(start_categorizations.values_list("includes_principles", "includes_draft")) == {
        (False, False),
        (False, True),
        (True, False),
        (True, True),
    }
    assert not AttributeCategorization.objects.filter(pk=stale.pk).exists()
